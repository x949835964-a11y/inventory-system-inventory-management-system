"""Excel reports with openpyxl.

Totals, values and stock status are real formulas, not pasted numbers: if the business owner
corrects a quantity or a price in the file, every total and status updates by itself.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.report_types import REASON_SUMMARY_ORDER, InventoryReport, MovementsReport

FONT = "Arial"
MONEY = '"S/" #,##0.00'
INTEGER = "#,##0"
DECIMAL3 = "#,##0.000"
DATETIME = "dd/mm/yyyy hh:mm"

HEADER_FILL = PatternFill("solid", start_color="1F2937")
TOTAL_FILL = PatternFill("solid", start_color="F3F4F6")
THIN = Side(style="thin", color="D1D5DB")
STATUS_FILLS = {"Agotado": "FDE2E1", "Stock bajo": "FEF3C7"}  # light red / light amber


def _font(**kwargs) -> Font:
    return Font(name=FONT, **kwargs)


def _title_block(ws: Worksheet, title: str, business: str, generated_at, filters: list[str]) -> None:
    ws["A1"] = title
    ws["A1"].font = _font(bold=True, size=14)
    ws["A2"] = business
    ws["A2"].font = _font(bold=True, size=11, color="374151")
    ws["A3"] = f"Generado el {generated_at:%d/%m/%Y %H:%M} · {' · '.join(filters)}"
    ws["A3"].font = _font(size=9, color="6B7280")


def _header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = _font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _widths(ws: Worksheet, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def _body_style(ws: Worksheet, first_row: int, last_row: int, last_col: int) -> None:
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, max_col=last_col):
        for cell in row:
            if cell.font.name != FONT:  # keep fonts we already set on purpose (e.g. colored type)
                cell.font = _font(size=10)
            cell.border = Border(bottom=THIN)


def _save(wb: Workbook) -> BytesIO:
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# --- Inventory -----------------------------------------------------------------------------------

INVENTORY_HEADERS = ["SKU", "Producto", "Categoría", "Proveedor", "Unidad", "Stock", "Stock mínimo",
                     "Precio compra (S/)", "Precio venta (S/)", "Valor en stock (S/)", "Estado"]


def inventory_xlsx(report: InventoryReport) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inventario"
    wb.properties.creator = report.business_name
    wb.properties.title = "Reporte de inventario"

    _title_block(ws, "Reporte de inventario actual", report.business_name, report.generated_at, report.filters)
    header = 5
    _header_row(ws, header, INVENTORY_HEADERS)
    first = header + 1

    if not report.rows:
        ws.cell(row=first, column=1, value="No hay productos para los filtros elegidos.").font = _font(italic=True)
        _widths(ws, [14, 34, 16, 24, 9, 11, 12, 14, 14, 16, 13])
        return _save(wb)

    for i, item in enumerate(report.rows):
        r = first + i
        qty_format = DECIMAL3 if item.allows_fractions else INTEGER
        values = [item.sku, item.name, item.category, item.supplier, item.unit,
                  item.stock, item.min_stock, item.purchase_price, item.sale_price]
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value)
        ws.cell(row=r, column=6).number_format = qty_format
        ws.cell(row=r, column=7).number_format = qty_format
        ws.cell(row=r, column=8).number_format = MONEY
        ws.cell(row=r, column=9).number_format = MONEY
        ws.cell(row=r, column=10, value=f"=F{r}*H{r}").number_format = MONEY
        ws.cell(row=r, column=11, value=f'=IF(F{r}<=0,"Agotado",IF(F{r}<=G{r},"Stock bajo","Disponible"))')
        ws.cell(row=r, column=11).alignment = Alignment(horizontal="center")
    last = first + len(report.rows) - 1
    _body_style(ws, first, last, len(INVENTORY_HEADERS))

    # The status column is a formula, so its color is a rule too (it follows edits).
    for status, color in STATUS_FILLS.items():
        ws.conditional_formatting.add(
            f"A{first}:K{last}",
            FormulaRule(formula=[f'$K{first}="{status}"'], fill=PatternFill("solid", start_color=color, end_color=color)),
        )

    total = last + 2
    summary = [
        ("Productos", f"=COUNTA(A{first}:A{last})", INTEGER),
        ("Productos en alerta", f'=COUNTIF(K{first}:K{last},"Stock bajo")+COUNTIF(K{first}:K{last},"Agotado")', INTEGER),
        ("Productos agotados", f'=COUNTIF(K{first}:K{last},"Agotado")', INTEGER),
        ("Valor total del inventario (S/)", f"=SUM(J{first}:J{last})", MONEY),
    ]
    for offset, (label, formula, fmt) in enumerate(summary):
        row = total + offset
        ws.cell(row=row, column=9, value=label).font = _font(bold=True)
        ws.cell(row=row, column=9).alignment = Alignment(horizontal="right")
        value = ws.cell(row=row, column=10, value=formula)
        value.font = _font(bold=True)
        value.number_format = fmt
        value.fill = TOTAL_FILL

    ws.freeze_panes = ws.cell(row=first, column=3)  # keep headers and SKU/name visible
    ws.auto_filter.ref = f"A{header}:K{last}"
    _widths(ws, [14, 34, 16, 24, 9, 11, 12, 14, 14, 16, 13])
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header}:{header}"
    return _save(wb)


# --- Movements -----------------------------------------------------------------------------------

MOVEMENT_HEADERS = ["Fecha", "SKU", "Producto", "Tipo", "Motivo", "Cantidad", "Unidad",
                    "Costo unit. (S/)", "Costo total (S/)", "Saldo", "Usuario", "Observaciones"]


def movements_xlsx(report: MovementsReport) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    wb.properties.creator = report.business_name
    wb.properties.title = "Reporte de movimientos"

    _title_block(ws, "Reporte de movimientos de inventario", report.business_name, report.generated_at, report.filters)
    header = 5
    _header_row(ws, header, MOVEMENT_HEADERS)
    first = header + 1
    _widths(ws, [16, 12, 30, 10, 22, 11, 8, 13, 14, 11, 12, 34])

    if not report.rows:
        ws.cell(row=first, column=1, value="No hay movimientos en el periodo elegido.").font = _font(italic=True)
        _summary_sheet(wb, report, first, first)  # zero totals, same layout
        return _save(wb)

    for i, m in enumerate(report.rows):
        r = first + i
        qty_format = DECIMAL3 if m.allows_fractions else INTEGER
        values = [m.moment, m.sku, m.product, m.type_label, m.reason, m.quantity, m.unit,
                  m.unit_cost, None, m.stock_after, m.user, m.notes]
        for col, value in enumerate(values, start=1):
            ws.cell(row=r, column=col, value=value)
        ws.cell(row=r, column=1).number_format = DATETIME
        ws.cell(row=r, column=6).number_format = qty_format
        ws.cell(row=r, column=8).number_format = MONEY
        ws.cell(row=r, column=9, value=f"=F{r}*H{r}").number_format = MONEY
        ws.cell(row=r, column=10).number_format = qty_format
        ws.cell(row=r, column=4).font = _font(color="1D4ED8" if m.is_entry else "C2410C", bold=True)
        ws.cell(row=r, column=12).alignment = Alignment(wrap_text=True, vertical="top")
    last = first + len(report.rows) - 1
    _body_style(ws, first, last, len(MOVEMENT_HEADERS))

    total = last + 2
    totals = [
        ("Total entradas (S/)", f'=SUMIF(D{first}:D{last},"Entrada",I{first}:I{last})'),
        ("Total salidas (S/)", f'=SUMIF(D{first}:D{last},"Salida",I{first}:I{last})'),
    ]
    for offset, (label, formula) in enumerate(totals):
        row = total + offset
        ws.cell(row=row, column=8, value=label).font = _font(bold=True)
        ws.cell(row=row, column=8).alignment = Alignment(horizontal="right")
        cell = ws.cell(row=row, column=9, value=formula)
        cell.font = _font(bold=True)
        cell.number_format = MONEY
        cell.fill = TOTAL_FILL

    ws.freeze_panes = ws.cell(row=first, column=1)
    ws.auto_filter.ref = f"A{header}:L{last}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header}:{header}"

    _summary_sheet(wb, report, first, last)
    return _save(wb)


def _summary_sheet(wb: Workbook, report: MovementsReport, first: int, last: int) -> None:
    """Totals per reason with COUNTIFS/SUMIFS over the detail sheet: always consistent with it."""
    ws = wb.create_sheet("Resumen")
    _title_block(ws, "Resumen por motivo", report.business_name, report.generated_at, report.filters)
    _header_row(ws, 5, ["Tipo", "Motivo", "N.º de movimientos", "Valor a costo (S/)"])
    _widths(ws, [12, 26, 18, 20])

    rng = lambda col: f"Movimientos!${col}${first}:${col}${last}"  # noqa: E731
    row = 6
    for type_label, reason in REASON_SUMMARY_ORDER:
        ws.cell(row=row, column=1, value=type_label)
        ws.cell(row=row, column=2, value=reason)
        criteria = f'{rng("D")},"{type_label}",{rng("E")},"{reason}"'
        ws.cell(row=row, column=3, value=f"=COUNTIFS({criteria})").number_format = INTEGER
        ws.cell(row=row, column=4, value=f"=SUMIFS({rng('I')},{criteria})").number_format = MONEY
        row += 1
    end = row - 1
    _body_style(ws, 6, end, 4)
    for col, value, fmt in ((2, "Total", None), (3, f"=SUM(C6:C{end})", INTEGER), (4, f"=SUM(D6:D{end})", MONEY)):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = _font(bold=True)
        cell.fill = TOTAL_FILL
        if fmt:
            cell.number_format = fmt
