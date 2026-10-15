"""Reports: data selection, Excel/PDF downloads and validation of the date range."""
from datetime import date, time
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.models import FRACTIONAL_UNITS as MODEL_FRACTIONAL_UNITS
from app.models import Category, MovementReason, MovementType
from app.services import report_service, stock_service
from app.services.exceptions import BusinessRuleError
from app.services.report_types import FRACTIONAL_UNITS
from app.utils.dates import local_to_utc

TZ = "America/Lima"


@pytest.fixture
def stock(db, category, make_product):
    tools = Category(name="Ferretería")
    db.session.add(tools)
    db.session.commit()
    return {
        "rice": make_product(sku="ARR-001", name="Arroz", stock=Decimal("50"), min_stock=Decimal("10")),
        "sugar": make_product(sku="AZU-001", name="Azúcar", stock=Decimal("3"), min_stock=Decimal("10"),
                              purchase_price=Decimal("4.00")),
        "hammer": make_product(sku="MAR-001", name="Martillo", category=tools, stock=Decimal("7"),
                               min_stock=Decimal("2"), purchase_price=Decimal("24.00")),
        "old": make_product(sku="VIE-001", name="Viejo", stock=Decimal("1"), is_active=False),
    }


@pytest.fixture
def history(admin, make_product):
    product = make_product(sku="DET-001", name="Detergente", stock=Decimal("0"), purchase_price=Decimal("10.00"))

    def move(movement_type, reason, qty, day, hour=12):
        stock_service.register_movement(
            product_id=product.id, movement_type=movement_type, reason=reason, quantity=Decimal(qty),
            user=admin, movement_date=local_to_utc(day, time(hour), TZ),
        )

    move(MovementType.IN, MovementReason.PURCHASE, "30", date(2025, 4, 1))
    move(MovementType.OUT, MovementReason.SALE, "4", date(2025, 4, 10))
    move(MovementType.OUT, MovementReason.SHRINKAGE, "1", date(2025, 4, 30), hour=22)  # 03:00 UTC on 01/05
    move(MovementType.OUT, MovementReason.SALE, "2", date(2025, 5, 2))
    return product


def test_fractional_units_match_the_model():
    # report_types keeps its own copy to stay free of database imports; keep them in sync.
    assert FRACTIONAL_UNITS == MODEL_FRACTIONAL_UNITS


# --- Inventory data -----------------------------------------------------------------------------

def test_inventory_report_lists_active_products_by_name(stock):
    report = report_service.inventory_report()
    assert [r.name for r in report.rows] == ["Arroz", "Azúcar", "Martillo"]
    assert report.total_value == Decimal("1105.00")  # 50*18.50 + 3*4 + 7*24
    assert report.low_stock_count == 1
    assert report.filters == ["Todos los productos activos"]


def test_inventory_report_filters(stock):
    tools = stock["hammer"].category_id
    by_category = report_service.inventory_report(category_id=tools)
    assert [r.name for r in by_category.rows] == ["Martillo"]
    assert by_category.filters == ["Categoría: Ferretería"]
    assert [r.name for r in report_service.inventory_report(stock="low").rows] == ["Azúcar"]


def test_inventory_report_unknown_category(stock):
    with pytest.raises(BusinessRuleError, match="no existe"):
        report_service.inventory_report(category_id=999)


# --- Movements data -----------------------------------------------------------------------------

def test_movements_report_uses_local_days_oldest_first(history):
    report = report_service.movements_report(date_from=date(2025, 4, 1), date_to=date(2025, 4, 30))
    assert [r.quantity for r in report.rows] == [Decimal("30"), Decimal("4"), Decimal("1")]  # 30/04 22:00 included
    assert (report.entry_count, report.exit_count) == (1, 2)
    assert (report.total_in_value, report.total_out_value) == (Decimal("300.00"), Decimal("50.00"))
    assert report.rows[2].moment.hour == 22  # shown in Lima time, not UTC


def test_movements_report_filters(history):
    exits = report_service.movements_report(date_from=date(2025, 4, 1), date_to=date(2025, 5, 31), movement_type="out")
    assert len(exits.rows) == 3 and exits.filters[-1] == "Solo salidas"
    none = report_service.movements_report(date_from=date(2025, 4, 1), date_to=date(2025, 5, 31), search="arroz")
    assert none.rows == []


def test_summary_by_reason(history):
    report = report_service.movements_report(date_from=date(2025, 4, 1), date_to=date(2025, 5, 31))
    summary = {(t, r): (c, v) for t, r, c, v in report.by_reason()}
    assert summary[("Salida", "Venta")] == (2, Decimal("60.00"))
    assert summary[("Salida", "Merma")] == (1, Decimal("10.00"))
    assert summary[("Entrada", "Ajuste de inventario")] == (0, Decimal("0.00"))


@pytest.mark.parametrize(
    ("date_from", "date_to", "message"),
    [(date(2025, 5, 1), date(2025, 4, 1), "posterior"), (date(2024, 1, 1), date(2025, 6, 1), "rango máximo")],
)
def test_invalid_ranges(db, date_from, date_to, message):
    with pytest.raises(BusinessRuleError, match=message):
        report_service.movements_report(date_from=date_from, date_to=date_to)


# --- Downloads -------------------------------------------------------------------------------------

def test_reports_page(stock, employee, client, login):
    login("empleado", "Empleado123")
    html = client.get("/reports/").get_data(as_text=True)
    assert "Inventario actual" in html and "Ferretería" in html
    assert "/reports/inventory.xlsx" in html and "/reports/movements.pdf" in html


def test_inventory_excel_download(stock, employee, client, login):
    login("empleado", "Empleado123")
    response = client.get("/reports/inventory.xlsx")
    assert response.status_code == 200
    assert response.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert 'attachment; filename=inventario_' in response.headers["Content-Disposition"]

    ws = load_workbook(BytesIO(response.data))["Inventario"]
    assert ws["A1"].value == "Reporte de inventario actual"
    assert [c.value for c in ws[5]][:3] == ["SKU", "Producto", "Categoría"]
    assert ws["A6"].value == "ARR-001"
    assert ws["J6"].value == "=F6*H6"                      # live formula, not a pasted number
    assert ws["K6"].value.startswith("=IF(F6<=0")
    # Summary block two rows below the 3 product rows (6-8): count, alerts, out of stock, value.
    assert ws["J10"].value == "=COUNTA(A6:A8)"
    assert ws["J13"].value == "=SUM(J6:J8)"


def test_inventory_pdf_download(stock, employee, client, login):
    login("empleado", "Empleado123")
    response = client.get("/reports/inventory.pdf?category=&stock=low")
    assert response.status_code == 200
    assert response.mimetype == "application/pdf"
    assert response.data.startswith(b"%PDF")


def test_movements_excel_download(history, admin, client, login):
    login("admin", "Admin123!")
    response = client.get("/reports/movements.xlsx?from=2025-04-01&to=2025-05-31")
    assert response.status_code == 200
    assert "movimientos_2025-04-01_a_2025-05-31.xlsx" in response.headers["Content-Disposition"]
    wb = load_workbook(BytesIO(response.data))
    assert wb.sheetnames == ["Movimientos", "Resumen"]
    ws = wb["Movimientos"]
    assert [ws.cell(r, 4).value for r in range(6, 10)] == ["Entrada", "Salida", "Salida", "Salida"]
    assert ws["I6"].value == "=F6*H6"
    assert wb["Resumen"]["C6"].value.startswith("=COUNTIFS(Movimientos!")


def test_movements_pdf_download(history, admin, client, login):
    login("admin", "Admin123!")
    response = client.get("/reports/movements.pdf?from=2025-04-01&to=2025-05-31&type=out")
    assert response.status_code == 200 and response.data.startswith(b"%PDF")


def test_empty_reports_still_download(db, admin, client, login):
    login("admin", "Admin123!")
    assert client.get("/reports/inventory.xlsx").status_code == 200
    assert client.get("/reports/movements.pdf?from=2020-01-01&to=2020-01-31").status_code == 200


def test_invalid_range_redirects_with_message(admin, client, login):
    login("admin", "Admin123!")
    response = client.get("/reports/movements.xlsx?from=2025-05-01&to=2025-04-01", follow_redirects=True)
    assert response.request.path == "/reports/"
    assert "posterior a la fecha final" in response.get_data(as_text=True)


def test_unknown_format_is_404(admin, client, login):
    login("admin", "Admin123!")
    assert client.get("/reports/inventory.csv").status_code == 404


def test_reports_require_login(client):
    response = client.get("/reports/inventory.xlsx")
    assert response.status_code == 302 and "/auth/login" in response.headers["Location"]
